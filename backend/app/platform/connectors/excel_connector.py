from openpyxl import load_workbook

from app.platform.connectors.contracts import CancellationToken, ColumnItem, DataBatch, PageRequest, Selection, TableItem
from app.platform.connectors.errors import ConnectorError, ConnectorErrorCode
from app.platform.connectors.file_base import FileConnector


class ExcelConnector(FileConnector):
    suffixes = {".xlsx"}

    def _sheet_name(self, selection: Selection) -> str | None:
        return selection.table or self.config.options.get("sheet")

    def _read(self, selection: Selection, page: PageRequest | None, cancellation: CancellationToken | None) -> tuple[list[str], list[dict], int]:
        self._check(cancellation)
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            sheet_name = self._sheet_name(selection)
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise ConnectorError(ConnectorErrorCode.DATA_ERROR, "工作表不存在")
            sheet = workbook[sheet_name] if sheet_name else workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
            if not headers or any(not header for header in headers) or len(headers) != len(set(headers)):
                raise ConnectorError(ConnectorErrorCode.DATA_ERROR, "Excel 表头为空或重复")
            all_rows = []
            for index, values in enumerate(iterator):
                if cancellation and index % 100 == 0:
                    cancellation.raise_if_cancelled()
                if any(value is not None for value in values):
                    all_rows.append(dict(zip(headers, values, strict=False)))
            rows = all_rows if page is None else all_rows[page.offset:page.offset + page.limit]
            return headers, rows, len(all_rows)
        finally:
            workbook.close()

    def discover_tables(self, schema: str | None = None, *, cancellation: CancellationToken | None = None) -> tuple[TableItem, ...]:
        del schema
        self._check(cancellation)
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            return tuple(TableItem(name=name, schema="file", kind="worksheet") for name in workbook.sheetnames)
        finally:
            workbook.close()

    def discover_columns(self, selection: Selection, *, cancellation: CancellationToken | None = None) -> tuple[ColumnItem, ...]:
        headers, rows, _ = self._read(selection, PageRequest(limit=100), cancellation)
        return self._columns_from_rows(headers, rows)

    def read_batch(self, selection: Selection, page: PageRequest = PageRequest(), *, cancellation: CancellationToken | None = None) -> DataBatch:
        headers, rows, total = self._read(selection, page, cancellation)
        return self._batch(self._columns_from_rows(headers, rows), rows, page, total)

    def total_rows(self, selection: Selection, *, cancellation: CancellationToken | None = None) -> int:
        return self._read(selection, None, cancellation)[2]

