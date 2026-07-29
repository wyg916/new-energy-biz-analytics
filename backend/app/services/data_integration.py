import csv
import hashlib
import json
import os
import re
import time
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlparse

import httpx
import psycopg
import pymysql
from openpyxl import load_workbook
from sqlalchemy import delete, func, select, text
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.auth import User
from app.models.integration import (
    DataIngestionQualityCheck, DataIngestionReview, DataIngestionRun,
    DataSetDefinition, DataSourceConnection, IngestedStationPreview,
    PublishedStationSnapshot,
)
from app.scenarios.charging_ops.manifest import METRICS, SCENARIO_ID, VERSION
from app.scenarios.registry import published_charging_ops
from app.services.dashboard import DashboardService
from app.services.metrics import MetricService

IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
STANDARD_FIELDS = {
    "station_id", "station_name", "region_id", "city_id", "charging_revenue",
    "charging_volume_kwh", "gross_profit", "gross_margin",
}
NUMERIC_FIELDS = {"charging_revenue", "charging_volume_kwh", "gross_profit", "gross_margin"}


class DataIntegrationError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _utc_now() -> datetime:
    return datetime.now(UTC)


def _safe_error_code(exc: Exception) -> str:
    if isinstance(exc, DataIntegrationError):
        return exc.code
    return f"CONNECTOR_{type(exc).__name__.upper()}"[:64]


def _decimal(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise DataIntegrationError("INVALID_NUMERIC_VALUE", "数值字段格式不合法") from exc


class DataIntegrationService:
    def __init__(self, db: Session, user: User | None = None):
        self.db = db
        self.user = user
        self.settings = get_settings()

    def overview(self, start: date, end_exclusive: date) -> dict:
        sources = list(self.db.scalars(select(DataSourceConnection).order_by(DataSourceConnection.created_at, DataSourceConnection.source_id)))
        dataset = self.db.get(DataSetDefinition, "station-operations")
        if dataset is None:
            raise DataIntegrationError("DATASET_NOT_CONFIGURED", "场站经营数据集尚未配置")
        mapping = json.loads(dataset.mapping_json)
        stored = list(self.db.scalars(
            select(IngestedStationPreview)
            .where(IngestedStationPreview.dataset_id == dataset.dataset_id)
            .order_by(IngestedStationPreview.id)
            .limit(5)
        ))
        if stored:
            preview = [self._preview_model(row) for row in stored]
            preview_batch_id = stored[0].batch_id
        else:
            if self.user is None:
                raise DataIntegrationError("USER_CONTEXT_REQUIRED", "读取业务预览需要用户权限上下文")
            station_rows = DashboardService(self.db, self.user).station_analysis(
                ["charging_revenue", "charging_volume_kwh", "gross_profit", "gross_margin"],
                start, end_exclusive, 5,
            )
            preview = [
                {
                    "station_id": row["station_id"],
                    "station_name": row["station_name"],
                    "region_id": row["region_id"],
                    "city_id": row["city_id"],
                    "charging_revenue": row["metrics"]["charging_revenue"],
                    "charging_volume_kwh": row["metrics"]["charging_volume_kwh"],
                    "gross_profit": row["metrics"]["gross_profit"],
                    "gross_margin": row["metrics"]["gross_margin"],
                }
                for row in station_rows["rows"]
            ]
            preview_batch_id = station_rows["metadata"]["batch_id"]
        latest_run = self.db.scalar(
            select(DataIngestionRun)
            .where(DataIngestionRun.dataset_id == dataset.dataset_id)
            .order_by(DataIngestionRun.started_at.desc(), DataIngestionRun.run_id.desc())
        )
        review = self.db.scalar(
            select(DataIngestionReview).where(DataIngestionReview.run_id == latest_run.run_id)
        ) if latest_run else None
        checks = list(self.db.scalars(
            select(DataIngestionQualityCheck)
            .where(DataIngestionQualityCheck.run_id == latest_run.run_id)
            .order_by(DataIngestionQualityCheck.rule_id)
        )) if latest_run else []
        source_payload = [self._source_payload(source) for source in sources]
        validations = {
            "mapping": all(item.get("source") and item.get("standard") in STANDARD_FIELDS for item in mapping),
            "sample": bool(preview),
            "types": all(row.get("charging_revenue") is None or isinstance(row.get("charging_revenue"), (int, float, Decimal)) for row in preview),
            "time_range": start < end_exclusive,
            "primary_key": len({row["station_id"] for row in preview}) == len(preview),
            "lineage": bool(preview_batch_id),
        }
        return {
            "sources": source_payload,
            "dataset": {
                "dataset_id": dataset.dataset_id,
                "display_name": dataset.display_name,
                "source_id": dataset.source_id,
                "source_object": dataset.source_object,
                "target_table": dataset.target_table,
                "standard_schema": dataset.standard_schema,
                "mapping": mapping,
                "status": dataset.status,
                "data_classification": dataset.data_classification,
            },
            "preview": preview,
            "validations": validations,
            "latest_ingestion": self._run_payload(latest_run) if latest_run else None,
            "workflow": self._review_payload(review, checks) if review else None,
            "metadata": {
                "data_classification": dataset.data_classification,
                "source": "platform_database",
                "batch_id": preview_batch_id,
                "data_time_range": {"start": start.isoformat(), "end_exclusive": end_exclusive.isoformat()},
                "generated_at": _utc_now().isoformat(),
            },
        }

    def test_source(self, source_id: str, password: str | None = None, resource_locator: str | None = None) -> dict:
        source = self.db.get(DataSourceConnection, source_id)
        if source is None:
            raise DataIntegrationError("SOURCE_NOT_FOUND", "数据源不存在")
        started = time.perf_counter()
        try:
            details = self._test_by_type(source, password, resource_locator)
            source.status = "available"
            source.last_error_code = None
            if resource_locator and source.source_type in {"excel", "api"}:
                source.resource_locator = resource_locator
            ok = True
        except Exception as exc:
            code = _safe_error_code(exc)
            source.status = "configured" if code in {"CREDENTIAL_REQUIRED", "RESOURCE_NOT_CONFIGURED"} else "unavailable"
            source.last_error_code = code
            source.last_tested_at = _utc_now()
            source.last_latency_ms = round((time.perf_counter() - started) * 1000)
            source.updated_at = _utc_now()
            self.db.commit()
            message = exc.message if isinstance(exc, DataIntegrationError) else "连接测试失败，详细异常未暴露"
            raise DataIntegrationError(code, message) from exc
        source.last_tested_at = _utc_now()
        source.last_latency_ms = round((time.perf_counter() - started) * 1000)
        source.updated_at = _utc_now()
        self.db.commit()
        return {
            "source_id": source.source_id,
            "source_type": source.source_type,
            "status": "passed" if ok else "failed",
            "latency_ms": source.last_latency_ms,
            "tested_at": source.last_tested_at.isoformat(),
            "details": details,
            "credential_persisted": False,
        }

    def ingest_dataset(self, dataset_id: str, start: date, end_exclusive: date, password: str | None = None, limit: int = 100) -> dict:
        dataset = self.db.get(DataSetDefinition, dataset_id)
        if dataset is None:
            raise DataIntegrationError("DATASET_NOT_FOUND", "数据集不存在")
        source = self.db.get(DataSourceConnection, dataset.source_id)
        if source is None:
            raise DataIntegrationError("SOURCE_NOT_FOUND", "数据源不存在")
        run = DataIngestionRun(
            run_id=f"ING-{uuid.uuid4()}",
            dataset_id=dataset.dataset_id,
            source_id=source.source_id,
            trigger_type="manual",
            status="running",
            started_at=_utc_now(),
        )
        self.db.add(run)
        self.db.commit()
        try:
            records = self._read_records(source, dataset, start, end_exclusive, password, limit)
            mapping = json.loads(dataset.mapping_json)
            normalized = [self._normalize_record(record, mapping) for record in records]
            batch_id = f"INGEST-{uuid.uuid4()}"
            checksum = hashlib.sha256(json.dumps(normalized, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()
            self.db.execute(delete(IngestedStationPreview).where(IngestedStationPreview.dataset_id == dataset.dataset_id))
            for index, record in enumerate(normalized):
                self.db.add(IngestedStationPreview(
                    dataset_id=dataset.dataset_id,
                    source_id=source.source_id,
                    source_record_id=str(record["station_id"] or index),
                    batch_id=batch_id,
                    data_classification=dataset.data_classification,
                    **record,
                ))
            run.rows_read = len(records)
            run.rows_written = len(normalized)
            run.rows_rejected = 0
            run.source_checksum = checksum
            run.status = "completed"
            run.finished_at = _utc_now()
            run.detail_json = json.dumps({
                "batch_id": batch_id, "target_table": dataset.target_table,
                "start": start.isoformat(), "end_exclusive": end_exclusive.isoformat(),
            })
            dataset.status = "validated"
            dataset.updated_at = _utc_now()
            self.db.commit()
            return self._run_payload(run)
        except Exception as exc:
            self.db.rollback()
            failed = self.db.get(DataIngestionRun, run.run_id)
            if failed:
                failed.status = "failed"
                failed.error_code = _safe_error_code(exc)
                failed.finished_at = _utc_now()
                self.db.commit()
            if isinstance(exc, DataIntegrationError):
                raise
            raise DataIntegrationError(_safe_error_code(exc), "接入任务失败，详细异常未暴露") from exc

    def validate_ingestion(self, dataset_id: str, run_id: str) -> dict:
        self._require_governance_role()
        dataset, run = self._require_current_completed_run(dataset_id, run_id)
        review = self.db.scalar(select(DataIngestionReview).where(DataIngestionReview.run_id == run.run_id))
        if review and review.workflow_status in {"pending_approval", "approved", "published"}:
            raise DataIntegrationError("WORKFLOW_INVALID_STATE", "已进入审批流程的批次不能重新执行质量校验")

        detail = json.loads(run.detail_json or "{}")
        batch_id = detail.get("batch_id")
        rows = list(self.db.scalars(
            select(IngestedStationPreview)
            .where(
                IngestedStationPreview.dataset_id == dataset.dataset_id,
                IngestedStationPreview.batch_id == batch_id,
            )
            .order_by(IngestedStationPreview.id)
        ))
        station_ids = [row.station_id for row in rows]
        required_complete = all(row.station_id and row.station_name and row.region_id for row in rows)
        ranges_valid = all(
            (row.charging_revenue is None or row.charging_revenue >= 0)
            and (row.charging_volume_kwh is None or row.charging_volume_kwh >= 0)
            and (row.gross_margin is None or Decimal("-1") <= row.gross_margin <= Decimal("1"))
            for row in rows
        )
        forbidden_tokens = {
            "name", "phone", "mobile", "id_card", "license_plate",
            "payment_account", "longitude", "latitude",
        }
        stored_columns = {column.name.lower() for column in IngestedStationPreview.__table__.columns}
        rule_specs = [
            ("DQI-001", "批次包含可发布数据", bool(rows), {"rows": len(rows)}),
            ("DQI-002", "关键字段完整", required_complete, {"required": ["station_id", "station_name", "region_id"]}),
            ("DQI-003", "场站主键唯一", len(station_ids) == len(set(station_ids)), {"distinct": len(set(station_ids))}),
            (
                "DQI-004", "接入行数对账",
                len(rows) == run.rows_written == run.rows_read and run.rows_rejected == 0,
                {"rows_read": run.rows_read, "rows_written": run.rows_written, "rows_rejected": run.rows_rejected},
            ),
            ("DQI-005", "数值范围符合业务规则", ranges_valid, {"revenue_and_energy_non_negative": True, "margin_range": [-1, 1]}),
            (
                "DQI-006", "批次溯源与分类完整",
                bool(batch_id) and all(row.data_classification == "simulated" for row in rows),
                {"batch_id": batch_id, "data_classification": dataset.data_classification},
            ),
            ("DQI-007", "源数据校验摘要存在", bool(run.source_checksum), {"checksum_algorithm": "sha256"}),
            (
                "DQI-008", "受控结构不含直接身份字段",
                not bool(forbidden_tokens & stored_columns),
                {"forbidden_fields_present": sorted(forbidden_tokens & stored_columns)},
            ),
        ]

        self.db.execute(delete(DataIngestionQualityCheck).where(DataIngestionQualityCheck.run_id == run.run_id))
        now = _utc_now()
        checks: list[DataIngestionQualityCheck] = []
        for rule_id, rule_name, passed, rule_detail in rule_specs:
            check = DataIngestionQualityCheck(
                check_id=f"DQC-{uuid.uuid4()}",
                run_id=run.run_id,
                dataset_id=dataset.dataset_id,
                rule_id=rule_id,
                rule_name=rule_name,
                severity="blocking",
                status="passed" if passed else "failed",
                detail_json=json.dumps(rule_detail, ensure_ascii=False),
                checked_at=now,
            )
            self.db.add(check)
            checks.append(check)
        failures = [check.rule_id for check in checks if check.status != "passed"]
        quality_status = "passed" if not failures else "failed"
        summary = {
            "rules_checked": len(checks),
            "rules_passed": len(checks) - len(failures),
            "failures": failures,
            "batch_id": batch_id,
        }
        if review is None:
            review = DataIngestionReview(
                review_id=f"REV-{uuid.uuid4()}",
                run_id=run.run_id,
                dataset_id=dataset.dataset_id,
            )
            self.db.add(review)
        review.quality_status = quality_status
        review.workflow_status = "quality_passed" if quality_status == "passed" else "quality_failed"
        review.quality_summary_json = json.dumps(summary, ensure_ascii=False)
        review.updated_at = now
        dataset.status = "quality_passed" if quality_status == "passed" else "quality_failed"
        dataset.updated_at = now
        self.db.commit()
        return self._review_payload(review, checks)

    def submit_ingestion(self, dataset_id: str, run_id: str) -> dict:
        self._require_governance_role()
        dataset, run = self._require_current_completed_run(dataset_id, run_id)
        review = self._require_review(run.run_id)
        if review.quality_status != "passed" or review.workflow_status != "quality_passed":
            raise DataIntegrationError("WORKFLOW_INVALID_STATE", "只有质量校验通过的批次可以提交审批")
        review.workflow_status = "pending_approval"
        review.requested_by = self.user.id
        review.requested_at = _utc_now()
        review.updated_at = review.requested_at
        dataset.status = "pending_approval"
        dataset.updated_at = review.requested_at
        self.db.commit()
        return self._review_payload(review, self._checks(run.run_id))

    def decide_ingestion(self, dataset_id: str, run_id: str, action: str, reason: str | None = None) -> dict:
        self._require_governance_role()
        dataset, run = self._require_current_completed_run(dataset_id, run_id)
        review = self._require_review(run.run_id)
        if review.workflow_status != "pending_approval":
            raise DataIntegrationError("WORKFLOW_INVALID_STATE", "当前批次不在待审批状态")
        if action not in {"approve", "reject"}:
            raise DataIntegrationError("INVALID_REVIEW_ACTION", "审批动作仅支持 approve 或 reject")
        if action == "reject" and not reason:
            raise DataIntegrationError("REJECTION_REASON_REQUIRED", "驳回审批必须填写原因")
        now = _utc_now()
        review.workflow_status = "approved" if action == "approve" else "rejected"
        review.decided_by = self.user.id
        review.decided_at = now
        review.rejection_reason = reason if action == "reject" else None
        review.updated_at = now
        dataset.status = review.workflow_status
        dataset.updated_at = now
        self.db.commit()
        return self._review_payload(review, self._checks(run.run_id))

    def publish_ingestion(self, dataset_id: str, run_id: str) -> dict:
        self._require_governance_role()
        dataset, run = self._require_current_completed_run(dataset_id, run_id)
        review = self._require_review(run.run_id)
        if review.workflow_status != "approved" or review.quality_status != "passed":
            raise DataIntegrationError("WORKFLOW_INVALID_STATE", "只有质量通过且审批通过的批次可以发布")
        published_count = self.db.scalar(
            select(func.count()).select_from(DataIngestionReview).where(
                DataIngestionReview.dataset_id == dataset.dataset_id,
                DataIngestionReview.workflow_status == "published",
            )
        ) or 0
        now = _utc_now()
        review.workflow_status = "published"
        review.release_version = f"v{published_count + 1}.0"
        review.published_by = self.user.id
        review.published_at = now
        review.updated_at = now
        dataset.status = "published"
        dataset.updated_at = now
        scenario = published_charging_ops(self.db)
        if scenario is None or not scenario.source_batch_id:
            raise DataIntegrationError("SCENARIO_NOT_PUBLISHED", "charging_ops 场景包尚未绑定已发布事实批次")
        detail = json.loads(run.detail_json or "{}")
        if not detail.get("start") or not detail.get("end_exclusive"):
            raise DataIntegrationError("INGESTION_PERIOD_MISSING", "接入批次缺少不可变快照所需的数据期间")
        period_start = date.fromisoformat(detail["start"])
        period_end = date.fromisoformat(detail["end_exclusive"])
        preview_rows = list(self.db.scalars(
            select(IngestedStationPreview).where(
                IngestedStationPreview.dataset_id == dataset.dataset_id,
                IngestedStationPreview.batch_id == detail.get("batch_id"),
            ).order_by(IngestedStationPreview.id)
        ))
        metric_service = MetricService(self.db)
        snapshot_digests: list[str] = []
        for row in preview_rows:
            if source := self.db.get(DataSourceConnection, run.source_id):
                if source.source_id == "platform-postgresql":
                    metrics = metric_service.compute(list(METRICS), period_start, period_end, [row.station_id])
                else:
                    metrics = {
                        "charging_revenue": float(row.charging_revenue) if row.charging_revenue is not None else None,
                        "charging_volume_kwh": float(row.charging_volume_kwh) if row.charging_volume_kwh is not None else None,
                        "gross_profit": float(row.gross_profit) if row.gross_profit is not None else None,
                        "gross_margin": float(row.gross_margin) if row.gross_margin is not None else None,
                    }
            payload = {
                "station_id": row.station_id, "station_name": row.station_name,
                "region_id": row.region_id, "city_id": row.city_id, "metrics": metrics,
            }
            digest = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
            snapshot_digests.append(digest)
            self.db.add(PublishedStationSnapshot(
                review_id=review.review_id, dataset_id=dataset.dataset_id, run_id=run.run_id,
                scenario_id=SCENARIO_ID, scenario_version=VERSION,
                release_version=review.release_version, source_record_id=row.source_record_id,
                station_id=row.station_id, station_name=row.station_name,
                region_id=row.region_id, city_id=row.city_id,
                metrics_json=json.dumps(metrics, ensure_ascii=False, sort_keys=True),
                period_start=period_start, period_end_exclusive=period_end,
                data_classification=dataset.data_classification, snapshot_checksum=digest,
                published_at=now,
            ))
        summary = json.loads(review.quality_summary_json or "{}")
        summary["snapshot_rows"] = len(preview_rows)
        summary["snapshot_checksum"] = hashlib.sha256("".join(snapshot_digests).encode()).hexdigest()
        summary["scenario_version"] = VERSION
        summary["source_batch_id"] = scenario.source_batch_id
        review.quality_summary_json = json.dumps(summary, ensure_ascii=False)
        self.db.commit()
        return self._review_payload(review, self._checks(run.run_id))

    def _require_governance_role(self) -> None:
        if self.user is None or self.user.role != "analyst_admin":
            raise DataIntegrationError("FORBIDDEN", "仅数据分析师/管理员可执行数据治理操作")

    def _require_current_completed_run(self, dataset_id: str, run_id: str) -> tuple[DataSetDefinition, DataIngestionRun]:
        dataset = self.db.get(DataSetDefinition, dataset_id)
        if dataset is None:
            raise DataIntegrationError("DATASET_NOT_FOUND", "数据集不存在")
        run = self.db.get(DataIngestionRun, run_id)
        if run is None or run.dataset_id != dataset.dataset_id:
            raise DataIntegrationError("INGESTION_RUN_NOT_FOUND", "接入批次不存在")
        if run.status != "completed":
            raise DataIntegrationError("WORKFLOW_INVALID_STATE", "只有已完成的接入批次可以进入质量与发布流程")
        latest = self.db.scalar(
            select(DataIngestionRun)
            .where(DataIngestionRun.dataset_id == dataset.dataset_id)
            .order_by(DataIngestionRun.started_at.desc(), DataIngestionRun.run_id.desc())
        )
        if latest is None or latest.run_id != run.run_id:
            raise DataIntegrationError("STALE_INGESTION_RUN", "旧接入批次不能覆盖或发布当前数据预览")
        return dataset, run

    def _require_review(self, run_id: str) -> DataIngestionReview:
        review = self.db.scalar(select(DataIngestionReview).where(DataIngestionReview.run_id == run_id))
        if review is None:
            raise DataIntegrationError("QUALITY_REVIEW_NOT_FOUND", "请先执行质量校验")
        return review

    def _checks(self, run_id: str) -> list[DataIngestionQualityCheck]:
        return list(self.db.scalars(
            select(DataIngestionQualityCheck)
            .where(DataIngestionQualityCheck.run_id == run_id)
            .order_by(DataIngestionQualityCheck.rule_id)
        ))

    def _test_by_type(self, source: DataSourceConnection, password: str | None, resource_locator: str | None) -> dict:
        if source.source_id == "platform-postgresql":
            database = self.db.scalar(text("SELECT current_database()")) if self.db.bind and self.db.bind.dialect.name == "postgresql" else "test"
            self.db.execute(text("SELECT 1"))
            return {"database": database, "mode": "managed_read_only"}
        if source.source_type == "postgresql":
            credential = self._credential(source, password)
            with psycopg.connect(
                host=source.host, port=source.port or 5432, dbname=source.database_name,
                user=source.username, password=credential, connect_timeout=5,
                options="-c default_transaction_read_only=on",
            ) as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT version(), current_database(), current_user")
                    version, database, username = cursor.fetchone()
            return {"database": database, "username": username, "version": version.split(",")[0], "mode": "read_only"}
        if source.source_type == "mysql":
            credential = self._credential(source, password)
            connection = pymysql.connect(
                host=source.host, port=source.port or 3306, database=source.database_name,
                user=source.username, password=credential, connect_timeout=5,
                read_timeout=5, write_timeout=5, charset="utf8mb4",
            )
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT VERSION(), CURRENT_USER(), DATABASE()")
                    version, username, database = cursor.fetchone()
            finally:
                connection.close()
            return {"database": database, "username": username, "version": version, "mode": "read_only_test"}
        if source.source_type == "excel":
            path = self._import_path(resource_locator or source.resource_locator)
            headers, row_count = self._inspect_spreadsheet(path)
            return {"file_name": path.name, "column_count": len(headers), "row_count": row_count, "headers": headers[:20]}
        if source.source_type == "api":
            url = resource_locator or source.resource_locator
            self._validate_api_url(url)
            with httpx.Client(timeout=5, follow_redirects=False, trust_env=False) as client:
                response = client.get(url)
                response.raise_for_status()
                payload = response.json()
            rows = payload.get("rows", []) if isinstance(payload, dict) else payload
            if not isinstance(rows, list):
                raise DataIntegrationError("API_INVALID_PAYLOAD", "API 返回结果必须是数组或包含 rows 数组")
            return {"status_code": response.status_code, "row_count": len(rows), "host": urlparse(url).hostname}
        raise DataIntegrationError("UNSUPPORTED_SOURCE_TYPE", "暂不支持该数据源类型")

    def _read_records(self, source: DataSourceConnection, dataset: DataSetDefinition, start: date, end_exclusive: date, password: str | None, limit: int) -> list[dict]:
        if source.source_id == "platform-postgresql":
            if self.user is None:
                raise DataIntegrationError("USER_CONTEXT_REQUIRED", "平台数据接入需要用户权限上下文")
            result = DashboardService(self.db, self.user).station_analysis(
                ["charging_revenue", "charging_volume_kwh", "gross_profit", "gross_margin"],
                start, end_exclusive, min(limit, 30),
            )
            return [
                {
                    "station_id": row["station_id"], "station_name": row["station_name"],
                    "region_id": row["region_id"], "city_id": row["city_id"],
                    **row["metrics"],
                }
                for row in result["rows"]
            ]
        if source.source_type == "excel":
            return self._spreadsheet_rows(self._import_path(source.resource_locator), limit)
        if source.source_type == "api":
            self._validate_api_url(source.resource_locator)
            with httpx.Client(timeout=10, follow_redirects=False, trust_env=False) as client:
                response = client.get(source.resource_locator)
                response.raise_for_status()
                payload = response.json()
            rows = payload.get("rows", []) if isinstance(payload, dict) else payload
            if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows[:limit]):
                raise DataIntegrationError("API_INVALID_PAYLOAD", "API 数据行格式不合法")
            return rows[:limit]
        return self._database_rows(source, dataset, password, limit)

    def _database_rows(self, source: DataSourceConnection, dataset: DataSetDefinition, password: str | None, limit: int) -> list[dict]:
        mapping = json.loads(dataset.mapping_json)
        columns = [field["source"] for field in mapping]
        if not dataset.source_object or not all(IDENTIFIER.fullmatch(part) for part in dataset.source_object.split(".")):
            raise DataIntegrationError("UNSAFE_SOURCE_OBJECT", "源表名称不符合安全规则")
        if not all(IDENTIFIER.fullmatch(column) for column in columns):
            raise DataIntegrationError("UNSAFE_SOURCE_COLUMN", "源字段名称不符合安全规则")
        credential = self._credential(source, password)
        if source.source_type == "postgresql":
            quoted_table = ".".join(f'"{part}"' for part in dataset.source_object.split("."))
            quoted_columns = ", ".join(f'"{column}"' for column in columns)
            sql = f"SELECT {quoted_columns} FROM {quoted_table} LIMIT %s"
            with psycopg.connect(
                host=source.host, port=source.port or 5432, dbname=source.database_name,
                user=source.username, password=credential, connect_timeout=5,
                options="-c default_transaction_read_only=on",
            ) as connection:
                with connection.cursor() as cursor:
                    cursor.execute(sql, (limit,))
                    return [dict(zip(columns, row, strict=True)) for row in cursor.fetchall()]
        if source.source_type == "mysql":
            quoted_table = ".".join(f"`{part}`" for part in dataset.source_object.split("."))
            sql = f'SELECT {", ".join(f"`{column}`" for column in columns)} FROM {quoted_table} LIMIT %s'
            connection = pymysql.connect(
                host=source.host, port=source.port or 3306, database=source.database_name,
                user=source.username, password=credential, connect_timeout=5,
                read_timeout=5, write_timeout=5, charset="utf8mb4",
                cursorclass=pymysql.cursors.DictCursor,
            )
            try:
                with connection.cursor() as cursor:
                    cursor.execute(sql, (limit,))
                    return list(cursor.fetchall())
            finally:
                connection.close()
        raise DataIntegrationError("UNSUPPORTED_DATABASE_SOURCE", "数据库类型不支持")

    def _normalize_record(self, record: dict, mapping: list[dict]) -> dict:
        normalized = {item["standard"]: record.get(item["source"]) for item in mapping if item.get("standard") in STANDARD_FIELDS}
        if not normalized.get("station_id") or not normalized.get("station_name") or not normalized.get("region_id"):
            raise DataIntegrationError("MISSING_REQUIRED_FIELD", "场站编码、名称和区域不能为空")
        for field in NUMERIC_FIELDS:
            normalized[field] = _decimal(normalized.get(field))
        normalized.setdefault("city_id", None)
        return {field: normalized.get(field) for field in STANDARD_FIELDS}

    def _credential(self, source: DataSourceConnection, supplied: str | None) -> str:
        credential = supplied or (os.getenv(source.credential_env_key) if source.credential_env_key else None)
        if not credential:
            raise DataIntegrationError("CREDENTIAL_REQUIRED", "该数据源需要单次凭据或受控环境变量")
        return credential

    def _import_path(self, locator: str | None) -> Path:
        if not locator:
            raise DataIntegrationError("RESOURCE_NOT_CONFIGURED", "尚未配置 Excel/CSV 文件")
        candidate = Path(locator)
        if candidate.is_absolute() or ".." in candidate.parts:
            raise DataIntegrationError("UNSAFE_FILE_PATH", "仅允许导入目录内的相对文件名")
        root = Path(self.settings.data_import_root).resolve()
        path = (root / candidate).resolve()
        if root not in path.parents and path != root:
            raise DataIntegrationError("UNSAFE_FILE_PATH", "文件路径超出受控导入目录")
        if not path.is_file():
            raise DataIntegrationError("FILE_NOT_FOUND", "导入文件不存在")
        if path.suffix.lower() not in {".xlsx", ".csv"}:
            raise DataIntegrationError("UNSUPPORTED_FILE_TYPE", "仅支持 xlsx 与 csv")
        return path

    def _inspect_spreadsheet(self, path: Path) -> tuple[list[str], int]:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.reader(handle)
                headers = next(reader, [])
                return headers, sum(1 for _ in reader)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value) for value in next(iterator, ()) if value is not None]
            return headers, sum(1 for row in iterator if any(value is not None for value in row))
        finally:
            workbook.close()

    def _spreadsheet_rows(self, path: Path, limit: int) -> list[dict]:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                return list(csv.DictReader(handle))[:limit]
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value) for value in next(rows, ())]
            return [dict(zip(headers, values, strict=False)) for _, values in zip(range(limit), rows, strict=False)]
        finally:
            workbook.close()

    def _validate_api_url(self, url: str | None) -> None:
        if not url:
            raise DataIntegrationError("RESOURCE_NOT_CONFIGURED", "尚未配置 API 地址")
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
            raise DataIntegrationError("UNSAFE_API_URL", "API 地址不符合安全规则")
        if parsed.hostname.lower() not in self.settings.api_source_allowed_hosts:
            raise DataIntegrationError("API_HOST_NOT_ALLOWED", "API 主机不在允许列表")

    @staticmethod
    def _preview_model(row: IngestedStationPreview) -> dict:
        return {
            "station_id": row.station_id, "station_name": row.station_name,
            "region_id": row.region_id, "city_id": row.city_id,
            "charging_revenue": float(row.charging_revenue) if row.charging_revenue is not None else None,
            "charging_volume_kwh": float(row.charging_volume_kwh) if row.charging_volume_kwh is not None else None,
            "gross_profit": float(row.gross_profit) if row.gross_profit is not None else None,
            "gross_margin": float(row.gross_margin) if row.gross_margin is not None else None,
        }

    @staticmethod
    def _source_payload(source: DataSourceConnection) -> dict:
        return {
            "source_id": source.source_id,
            "display_name": source.display_name,
            "source_type": source.source_type,
            "endpoint": {
                "host": source.host, "port": source.port,
                "database_name": source.database_name, "username": source.username,
                "resource_locator": source.resource_locator,
            },
            "status": source.status,
            "last_tested_at": source.last_tested_at.isoformat() if source.last_tested_at else None,
            "last_latency_ms": source.last_latency_ms,
            "last_error_code": source.last_error_code,
            "credential_stored": False,
        }

    @staticmethod
    def _run_payload(run: DataIngestionRun) -> dict:
        return {
            "run_id": run.run_id, "dataset_id": run.dataset_id, "source_id": run.source_id,
            "status": run.status, "rows_read": run.rows_read, "rows_written": run.rows_written,
            "rows_rejected": run.rows_rejected, "source_checksum": run.source_checksum,
            "error_code": run.error_code,
            "started_at": run.started_at.isoformat(),
            "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        }

    @staticmethod
    def _review_payload(
        review: DataIngestionReview,
        checks: list[DataIngestionQualityCheck],
    ) -> dict:
        summary = json.loads(review.quality_summary_json or "{}")
        return {
            "review_id": review.review_id,
            "run_id": review.run_id,
            "dataset_id": review.dataset_id,
            "quality_status": review.quality_status,
            "workflow_status": review.workflow_status,
            "release_version": review.release_version,
            "requested_by": review.requested_by,
            "requested_at": review.requested_at.isoformat() if review.requested_at else None,
            "decided_by": review.decided_by,
            "decided_at": review.decided_at.isoformat() if review.decided_at else None,
            "published_by": review.published_by,
            "published_at": review.published_at.isoformat() if review.published_at else None,
            "rejection_reason": review.rejection_reason,
            "summary": summary,
            "checks": [
                {
                    "rule_id": check.rule_id,
                    "rule_name": check.rule_name,
                    "severity": check.severity,
                    "status": check.status,
                    "detail": json.loads(check.detail_json or "{}"),
                    "checked_at": check.checked_at.isoformat(),
                }
                for check in checks
            ],
        }
