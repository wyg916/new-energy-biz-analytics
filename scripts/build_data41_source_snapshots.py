from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook


UCI_COLUMNS = (
    "source_row_number",
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build the committed DATA-4.1 snapshots from downloaded official sources."
    )
    parser.add_argument("--acn", type=Path, required=True)
    parser.add_argument("--uci-zip", type=Path, required=True)
    parser.add_argument("--uci-xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    args.output.mkdir(parents=True, exist_ok=True)
    acn_target = args.output / "acn_data_ornl_26_v1.json"
    uci_target = args.output / "uci_online_retail_every_20th_v1.csv"

    acn_payload = json.loads(args.acn.read_text(encoding="utf-8"))
    if acn_payload.get("total_count") != 26 or len(acn_payload.get("results", [])) != 26:
        raise RuntimeError("The ACN discovery snapshot must contain exactly 26 source records.")
    acn_target.write_text(
        json.dumps(acn_payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )

    workbook = load_workbook(args.uci_xlsx, read_only=True, data_only=True)
    sheet = workbook.active
    source_rows = 0
    selected_rows = 0
    source_min = None
    source_max = None
    selected_min = None
    selected_max = None
    with uci_target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(UCI_COLUMNS)
        for source_row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if source_row_number == 1:
                if tuple(row) != UCI_COLUMNS[1:]:
                    raise RuntimeError(f"Unexpected UCI header: {row!r}")
                continue
            source_rows += 1
            invoice_date = row[4]
            source_min = invoice_date if source_min is None else min(source_min, invoice_date)
            source_max = invoice_date if source_max is None else max(source_max, invoice_date)
            if source_rows % 20 != 0:
                continue
            selected_rows += 1
            selected_min = invoice_date if selected_min is None else min(selected_min, invoice_date)
            selected_max = invoice_date if selected_max is None else max(selected_max, invoice_date)
            customer_id = "" if row[6] is None else str(int(row[6]))
            writer.writerow(
                (
                    source_row_number,
                    str(row[0]),
                    str(row[1]),
                    "" if row[2] is None else str(row[2]),
                    int(row[3]),
                    invoice_date.isoformat(sep=" "),
                    str(row[5]),
                    customer_id,
                    "" if row[7] is None else str(row[7]),
                )
            )
    workbook.close()

    expected_rows = 541_909
    if source_rows != expected_rows or selected_rows != expected_rows // 20:
        raise RuntimeError(
            f"Unexpected UCI counts: source={source_rows}, selected={selected_rows}."
        )

    manifest = {
        "manifest_version": "data41-source-manifest-v1",
        "generated_at": datetime.now(UTC).isoformat(),
        "snapshots": [
            {
                "dataset_code": "charging_ops_acn_ornl_discovery",
                "source_name": "ACN-Data via ORNL OpenEnergyDataPortal",
                "source_url": "https://openenergyhub.ornl.gov/explore/dataset/acn-data/",
                "source_api_url": "https://openenergyhub.ornl.gov/api/explore/v2.1/catalog/datasets/acn-data/records?limit=100",
                "license": "CC BY 4.0",
                "publisher": "California Institute of Technology",
                "distributor": "Oak Ridge National Laboratory OpenEnergyDataPortal",
                "source_download_sha256": sha256_file(args.acn),
                "snapshot_path": "data/open_source/raw/acn_data_ornl_26_v1.json",
                "snapshot_sha256": sha256_file(acn_target),
                "source_row_count": 26,
                "snapshot_row_count": 26,
                "selection_method": "All 26 records exposed by the ORNL discovery copy.",
                "dataset_version": "ornl-acn-discovery-2020-v1",
                "transformation_version": "data41-acn-transform-v1",
                "run_id": "DATA41-ACN-ORNL-26-V1",
            },
            {
                "dataset_code": "sales_ops_uci_online_retail",
                "source_name": "UCI Machine Learning Repository - Online Retail",
                "source_url": "https://archive.ics.uci.edu/dataset/352/online+retail",
                "source_download_url": "https://archive.ics.uci.edu/static/public/352/online+retail.zip",
                "license": "CC BY 4.0",
                "publisher": "UCI Machine Learning Repository",
                "source_download_sha256": sha256_file(args.uci_zip),
                "snapshot_path": "data/open_source/raw/uci_online_retail_every_20th_v1.csv",
                "snapshot_sha256": sha256_file(uci_target),
                "source_row_count": source_rows,
                "snapshot_row_count": selected_rows,
                "source_period_start": source_min.isoformat(),
                "source_period_end": source_max.isoformat(),
                "snapshot_period_start": selected_min.isoformat(),
                "snapshot_period_end": selected_max.isoformat(),
                "selection_method": "Every twentieth data row in original workbook order; source_row_number is retained.",
                "dataset_version": "uci-352-online-retail-v1-sample-20",
                "transformation_version": "data41-uci-transform-v1",
                "run_id": "DATA41-UCI-ONLINE-RETAIL-352-V1",
            },
        ],
    }
    (args.output.parent / "source_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
